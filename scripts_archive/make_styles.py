with open('scripts_archive/report_styles.py', 'w', encoding='utf-8') as f:
    f.write(r'''#!/usr/bin/env python3
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

# Color Palette
COL_NAVY = RGBColor(0x1E, 0x3A, 0x8A)      # #1E3A8A
COL_SLATE = RGBColor(0x33, 0x41, 0x55)     # #334155
COL_SUBTLE = RGBColor(0x47, 0x55, 0x69)   # #475569
COL_BODY = RGBColor(0x1E, 0x29, 0x3B)      # #1E293B

# Table border colors
BORDER_OUTER = "334155"
BORDER_HEADER = "142651"
BORDER_HORIZ = "CBD5E1"
BORDER_VERT = "E2E8F0"
HEADER_BG = "F1F5F9"
ZEBRA_BG = "F8FAFC"

def setup_document():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)
        s.page_width = Inches(8.5)
        s.page_height = Inches(11.0)
    return doc

def add_title(doc, title_text, subtitle_text=None):
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.keep_with_next = True
    run = p_title.add_run(title_text)
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = COL_NAVY
    run.font.name = 'Calibri'

    if subtitle_text:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sub.paragraph_format.space_after = Pt(10)
        p_sub.paragraph_format.keep_with_next = True
        run_sub = p_sub.add_run(subtitle_text)
        run_sub.font.size = Pt(11)
        run_sub.font.italic = True
        run_sub.font.color.rgb = COL_SLATE
        run_sub.font.name = 'Calibri'

def add_meta_box(doc, course_str, group_str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    
    r1 = p.add_run(course_str + "\n")
    r1.font.size = Pt(9.5)
    r1.font.bold = True
    r1.font.color.rgb = COL_NAVY
    r1.font.name = 'Calibri'
    
    r2 = p.add_run(group_str)
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = COL_SLATE
    r2.font.name = 'Calibri'

def add_h1(doc, text):
    h = doc.add_heading('', level=1)
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = COL_NAVY
    run.font.name = 'Calibri'
    return h

def add_h2(doc, text):
    h = doc.add_heading('', level=2)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    run = h.add_run(text)
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.color.rgb = COL_SLATE
    run.font.name = 'Calibri'
    return h

def add_p(doc, text, space_after=4.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.size = Pt(10.5)
    run.font.color.rgb = COL_BODY
    run.font.name = 'Calibri'
    return p

def add_bullet(doc, bold_prefix, text=""):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(3.0)
    p.paragraph_format.line_spacing = 1.15
    
    r0 = p.add_run("\u2022  ")
    r0.font.size = Pt(10.5)
    r0.font.bold = True
    r0.font.color.rgb = COL_NAVY
    
    if bold_prefix:
        r1 = p.add_run(bold_prefix + " ")
        r1.font.size = Pt(10.0)
        r1.font.bold = True
        r1.font.color.rgb = COL_SLATE
        r1.font.name = 'Calibri'
    
    if text:
        r2 = p.add_run(text)
        r2.font.size = Pt(10.0)
        r2.font.color.rgb = COL_BODY
        r2.font.name = 'Calibri'
    return p

def add_figure(doc, img_path, caption):
    if not os.path.exists(img_path):
        alt_path = os.path.join('figures', img_path)
        if os.path.exists(alt_path):
            img_path = alt_path
        else:
            print(f"WARNING: Figure not found: {img_path}")
            return
    
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(8)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.paragraph_format.keep_with_next = True
    run = p_img.add_run()
    run.add_picture(img_path, width=Inches(5.95))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    run_cap = p_cap.add_run(caption)
    run_cap.font.size = Pt(9.0)
    run_cap.font.italic = True
    run_cap.font.color.rgb = COL_SUBTLE
    run_cap.font.name = 'Calibri'

def add_styled_table(doc, title, headers, data_rows, col_widths, col_alignments=None, footnote=None):
    p_til = doc.add_paragraph()
    p_til.paragraph_format.space_before = Pt(10)
    p_til.paragraph_format.space_after = Pt(3)
    p_til.paragraph_format.keep_with_next = True
    run_t = p_til.add_run(title)
    run_t.font.size = Pt(9.5)
    run_t.font.bold = True
    run_t.font.color.rgb = COL_NAVY
    run_t.font.name = 'Calibri'
    
    table = doc.add_table(rows=len(data_rows)+1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    
    borders = parse_xml(f"""
        <w:tblBorders {gnsdecls('w')}>
            <w:top w:val="single" w:sz="8" w:space="0" w:color="{BORDER_OUTER}"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{BORDER_OUTER}"/>
            <w:left w:val="single" w:sz="6" w:space="0" w:color="{BORDER_OUTER}"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="{BORDER_OUTER}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{BORDER_HORIZ}"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{BORDER_VERT}"/>
        </w:tblBorders>
    """)
    tblPr.append(borders)
    
    mar = parse_xml(f"""
        <w:tblCellMar {gnsdecls('w')}>
            <w:top w:w="120" w:type="dxa"/>
            <w:bottom w:w="120" w:type="dxa"/>
            <w:left w:w="140" w:type="dxa"/>
            <w:right w:w="140" w:type="dxa"/>
        </w:tblCellMar>
    """)
    tblPr.append(mar)
    
    hdr_row = table.rows[0]
    hdr_trPr = hdr_row._tr.get_or_add_trPr()
    hdr_trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    hdr_trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
    
    for i, h in enumerate(headers):
        w_dxa = int(col_widths[i] * 1440)
        hdr_row.cells[i].width = Inches(col_widths[i])
        tcPr = hdr_row.cells[i]._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="{w_dxa}" w:type="dxa"/>'))
        tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEADER_BG}"/>'))
        tcBorders = parse_xml(f"""
            <w:tcBorders {nsdecls('w')}>
                <w:bottom w:val="single" w:sz="10" w:space="0" w:color="{BORDER_HEADER}"/>
            </w:tcBorders>
        """)
        tcPr.append(tcBorders)
        
        p = hdr_row.cells[i].paragraphs[0]
        p.text = h
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2.5)
        p.paragraph_format.space_after = Pt(2.5)
        if p.runs:
            p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(8.5)
            p.runs[0].font.color.rgb = COL_NAVY
            p.runs[0].font.name = 'Calibri'
    
    for r1, row in enumerate(data_rows):
        r_idx = r1 + 1
        row_obj = table.rows[r_idx]
        row_trPr = row_obj._tr.get_or_add_trPr()
        row_trPr.append(parse_xml(f'<w:cantSplit {gnsdecls("w")}/>'))
        is_even = (r1 % 2 == 1)
        
        for c_idx, val in enumerate(row):
            w_dxa = int(col_widths[c_idx] * 1440)
            row_obj.cells[c_idx].width = Inches(col_widths[c_idx])
            tcPr = row_obj.cells[c_idx]._tc.get_or_add_tcPr()
            tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="{w_dxa}" w:type="dxa"/>'))
            if is_even:
                tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{ZEBRA_BG}"/>'))
            
            p = row_obj.cells[c_idx].paragraphs[0]
            p.text = str(val)
            if col_alignments and c_idx < len(col_alignments):
                p.alignment = col_alignments[c_idx]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WG_ALIGN_PARAGRAPH.CENTER
            
            p.paragraph_format.space_before = Pt(2.0)
            p.paragraph_format.space_after = Pt(2.0)
            if p.runs:
                p.runs[0].font.size = Pt(8.0)
                p.runs[0].font.color.rgb = COL_BODY
                p.runs[0].font.name = 'Calibri'
    
    if footnote:
        p_fn = doc.add_paragraph()
        p_fn.paragraph_format.space_before = Pt(1)
        p_fn.paragraph_format.space_after = Pt(8)
        run_fn = p_fn.add_run(footnote)
        run_fn.font.size = Pt(8.0)
        run_fn.font.italic = True
        run_fn.font.color.rgb = COL_SUBTLE
        run_fn.font.name = 'Calibri'
    else:
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_after = Pt(6)
    return table
''')
